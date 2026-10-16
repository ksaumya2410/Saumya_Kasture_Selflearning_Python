
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# Create a new Excel workbook
workbook = Workbook()

# Set the active worksheet title
sheet = workbook.active
sheet.title = "Semester_Results"

# Define the header row
headers = [
    "Roll No",
    "Student Name",
    "Lab Marks (out of 50)",
    "Theory Marks (out of 50)",
    "Total Marks (out of 100)",
    "Result"
]

sheet.append(headers)

# Sample academic records (at least 5 students)
students = [
    ("AIML001", "Aarav", 45, 42),
    ("AIML002", "Diya", 38, 35),
    ("AIML003", "Rohan", 20, 25),
    ("AIML004", "Ananya", 48, 46),
    ("AIML005", "Kabir", 30, 32)
]

total_scores = []

# Calculate total marks and result for each student
for roll_no, name, lab_marks, theory_marks in students:

    total = lab_marks + theory_marks

    if total >= 40:
        result = "Pass"
    else:
        result = "Fail"

    sheet.append([
        roll_no,
        name,
        lab_marks,
        theory_marks,
        total,
        result
    ])

    total_scores.append(total)


# Calculate the class average of total marks
class_average = sum(total_scores) / len(total_scores)

# Add the final summary row
sheet.append([
    "",
    "Class Average",
    "",
    "",
    round(class_average, 2),
    ""
])


# Style the header row
header_fill = PatternFill(
    fill_type="solid",
    fgColor="D9D9D9"
)

for cell in sheet[1]:
    cell.font = Font(bold=True)
    cell.fill = header_fill
    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True
    )


# Style the class average row
summary_row = sheet.max_row

for cell in sheet[summary_row]:
    cell.font = Font(bold=True)
    cell.fill = PatternFill(
        fill_type="solid",
        fgColor="EAEAEA"
    )


# Center the data in the worksheet
for row in sheet.iter_rows():
    for cell in row:
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

# Set suitable column widths
column_widths = {
    "A": 14,
    "B": 20,
    "C": 24,
    "D": 28,
    "E": 25,
    "F": 14
}

for column, width in column_widths.items():
    sheet.column_dimensions[column].width = width

# Adjust header row height
sheet.row_dimensions[1].height = 35

# Save the Excel workbook
filename = "AIML_Batch_Grades.xlsx"
workbook.save(filename)

print("Excel workbook created successfully!")
print("File name:", filename)
print("Number of student records:", len(students))
print(f"Class average total marks: {class_average:.2f}")
